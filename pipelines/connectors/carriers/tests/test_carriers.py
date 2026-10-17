"""Connector hãng vận chuyển (FR-ING-09): GHN, GHTK qua API giả lập (respx), J&T qua file đối soát."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import httpx
import openpyxl
import pytest
import respx

from connectors.base import SourceError, SyncContext
from connectors.carriers import ghn, ghtk, jnt
from dx_common.tenants import SourceConfig, TenantConfig
from dx_pipelines.warehouse import DuckDBWarehouse

NOW = datetime(2026, 10, 16, 3, 0, tzinfo=UTC)


def _ctx(tmp_path: Path, **extra: Any) -> SyncContext:
    wh = DuckDBWarehouse(tmp_path / "dx.duckdb")
    return SyncContext(wh, "t_demo_seller_raw", "run-1", NOW, extra={"dlt_dir": str(tmp_path / "dlt"), **extra})


def _tenant(*sources: SourceConfig) -> TenantConfig:
    return TenantConfig.model_validate(
        {
            "tenant_id": "demo_seller",
            "name": "Shop Demo",
            "pack": "seller",
            "sources": [s.model_dump() for s in sources],
        }
    )


def _event(waybill: str, status: str) -> dict[str, Any]:
    t = datetime(2026, 10, 10, 2, 0, tzinfo=UTC)
    return {
        "client_order_code": "O-" + waybill,
        "status_raw": status,
        "event_at": t,
        "first_at": t,
        "shipping_fee": Decimal(20000),
        "cod_amount": Decimal(300000),
    }


def _rows(ctx: SyncContext, table: str) -> list[tuple[Any, ...]]:
    return ctx.warehouse.query(
        f'select waybill, status_raw, shipping_fee, cod_amount, return_fee from "t_demo_seller_raw"."{table}" order by 1'
    )


def test_fr_ing_09_ghn_detail_refreshes_open_waybills(tmp_path: Path) -> None:
    src = SourceConfig(type="ghn", connector="dlt", account_id="2509001", secret_ref="dx-demo_seller-ghn-2509001")
    calls: list[str] = []

    def detail(request: httpx.Request) -> httpx.Response:
        assert request.headers["Token"] == "tok" and request.headers["ShopId"] == "2509001"
        code = __import__("json").loads(request.content)["order_code"]
        calls.append(code)
        return httpx.Response(
            200,
            json={
                "code": 200,
                "data": {
                    "order_code": code,
                    "client_order_code": "O-" + code,
                    "status": "returned",
                    "cod_amount": 300000,
                    "total_fee": 22000,
                    "fee": {"return": 11000, "insurance": 0},
                    "weight": 500,
                    "to_name": "Nguyễn Văn A",
                    "to_phone": "0909123456",
                    "updated_date": "2026-10-12T05:00:00Z",
                },
            },
        )

    with respx.mock() as m:
        m.post(f"{ghn.BASE}/shiip/public-api/v2/shipping-order/detail").mock(side_effect=detail)
        client = ghn.GhnClient("tok", "2509001", sleep=lambda s: None)
        events = {"W1": _event("W1", "delivering"), "W2": _event("W2", "delivered")}
        ctx = _ctx(tmp_path, ghn_client=client, ghn_events=events)
        out = ghn.sync_ghn(_tenant(src), src, ctx)
    assert calls == ["W1"]  # vận đơn đã ở trạng thái cuối không gọi lại API
    assert out.tables == {"ghn__shipments": 2}
    rows = _rows(ctx, "ghn__shipments")
    assert rows[0][:2] == ("W1", "returned") and rows[0][2] == 22000 and rows[0][4] == 11000
    assert rows[1][:2] == ("W2", "delivered") and rows[1][2] == 20000
    cols = [
        r[0]
        for r in ctx.warehouse.query(
            "select column_name from information_schema.columns where table_name = 'ghn__shipments'"
        )
    ]
    assert not {"to_name", "to_phone"} & set(cols)


def test_fr_ing_09_ghtk_auth_error(tmp_path: Path) -> None:
    with respx.mock() as m:
        m.get(f"{ghtk.BASE}/services/shipment/v2/W1").mock(
            return_value=httpx.Response(200, json={"success": False, "message": "Token không hợp lệ"})
        )
        client = ghtk.GhtkClient("bad", "S1", sleep=lambda s: None)
        with pytest.raises(SourceError) as exc:
            client.status("W1")
    assert exc.value.reason == "auth_revoked"


def test_fr_ing_09_ghtk_status_fields(tmp_path: Path) -> None:
    with respx.mock() as m:
        m.get(f"{ghtk.BASE}/services/shipment/v2/W9").mock(
            return_value=httpx.Response(
                200,
                json={
                    "success": True,
                    "order": {
                        "label_id": "W9",
                        "partner_id": "O9",
                        "status": 5,
                        "modified": "2026-10-12 10:00:00",
                        "ship_money": 30000,
                        "pick_money": 450000,
                        "insurance": 2000,
                        "customer_fullname": "B",
                        "customer_tel": "0911111111",
                    },
                },
            )
        )
        d = ghtk.GhtkClient("tok", "S1", sleep=lambda s: None).status("W9")
    assert d is not None and d["status_raw"] == "5" and d["cod_amount"] == 450000 and "customer_tel" not in d


def _xlsx(rows: list[list[Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEAD = [
    "Mã vận đơn",
    "Mã đơn hàng",
    "Trạng thái",
    "Phí vận chuyển",
    "Tiền COD",
    "Ngày gửi",
    "Phí hoàn",
    "Tên người nhận",
    "SĐT người nhận",
]


def test_fr_ing_09_jnt_reconcile_file(tmp_path: Path) -> None:
    good = _xlsx(
        [
            HEAD,
            ["JT1", "O1", "Đã ký nhận", "25.000", "350.000", "10/10/2026 08:00", 0, "Lan", "0909"],
            ["JT2", "O2", "Đã hoàn", 18000, 0, "11/10/2026 09:30", 9000, "Minh", "0912"],
        ]
    )
    bad = _xlsx([["Cột lạ", "Khác"], [1, 2]])
    src = SourceConfig(type="jnt", connector="dlt", account_id="file")
    files = [{"id": "f1", "filename": "jt.xlsx", "content": good}, {"id": "f2", "filename": "x.xlsx", "content": bad}]
    ctx = _ctx(tmp_path, jnt_files=files)
    out = jnt.sync_jnt(_tenant(src), src, ctx)
    assert out.tables == {"jnt__shipments": 2}
    rows = _rows(ctx, "jnt__shipments")
    assert rows == [("JT1", "Đã ký nhận", 25000, 350000, 0), ("JT2", "Đã hoàn", 18000, 0, 9000)]
    parsed, err = jnt.parse_file(bad, "x.xlsx")
    assert parsed == [] and err and "J&T" in err
