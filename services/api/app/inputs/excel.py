"""File Excel giá vốn (FR-INP-02, spec 05 §6.4).

Mẫu: sheet `gia_von` (dòng 1 là tiêu đề `ma_sku`, `gia_von`, `ap_dung_tu`, `ghi_chu`) và sheet `huong_dan`.
Xử lý: tối đa 5.000 dòng, 2 MB; `openpyxl` chế độ chỉ đọc, chỉ lấy giá trị (`data_only`). Dòng lỗi không làm hỏng
dòng khác. Không lưu nội dung file.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from dx_common.money import parse_vnd_text

MAX_ROWS = 5000
MAX_BYTES = 2 * 1024 * 1024
COLUMNS = ("ma_sku", "gia_von", "ap_dung_tu", "ghi_chu")

MESSAGES = {
    "REQUIRED": "Thiếu giá trị bắt buộc",
    "INVALID_NUMBER": "Giá vốn không phải số hợp lệ",
    "NEGATIVE": "Giá vốn phải lớn hơn hoặc bằng 0",
    "INVALID_DATE": "Ngày áp dụng phải theo dạng dd/mm/yyyy",
    "DUPLICATE_IN_FILE": "Trùng SKU và ngày áp dụng với một dòng khác trong file",
    "UNKNOWN_SKU": "SKU chưa thấy trong dữ liệu bán hàng (dòng vẫn được nhận)",
}


@dataclass
class RowError:
    row: int
    column: str
    code: str
    message_vi: str
    warning: bool = False

    def as_dict(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "column": self.column,
            "code": self.code,
            "message_vi": self.message_vi,
            "warning": self.warning,
        }


@dataclass
class ParsedRow:
    row: int
    seller_sku: str
    unit_cost_vnd: Decimal
    valid_from: date
    note: str | None


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    total: int = 0

    @property
    def error_rows(self) -> set[int]:
        return {e.row for e in self.errors if not e.warning}


class FileRejected(ValueError):
    def __init__(self, code: str, message_vi: str) -> None:
        super().__init__(code)
        self.code = code
        self.message_vi = message_vi


def _err(row: int, column: str, code: str) -> RowError:
    return RowError(row, column, code, MESSAGES[code], warning=code == "UNKNOWN_SKU")


def _parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v is None:
        return None
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_cost(v: Any) -> Decimal | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    try:
        return parse_vnd_text(str(v))
    except ValueError:
        return None


def _read_rows(content: bytes, filename: str) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = None
        for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise FileRejected("INVALID_FILE", "Không đọc được file CSV")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        return [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    if name.endswith(".xlsx"):
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise FileRejected("INVALID_FILE", "Không đọc được file Excel") from exc
        ws = wb["gia_von"] if "gia_von" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise FileRejected("INVALID_FILE_TYPE", "Chỉ nhận file .xlsx hoặc .csv")


def parse_costs_file(content: bytes, filename: str, known_skus: set[str] | None = None) -> ParseResult:
    if len(content) > MAX_BYTES:
        raise FileRejected("FILE_TOO_LARGE", "File vượt quá 2 MB")
    rows = _read_rows(content, filename)
    if not rows:
        raise FileRejected("EMPTY_FILE", "File trống")
    header = [str(h or "").strip().lower() for h in rows[0]]
    missing = [c for c in COLUMNS[:3] if c not in header]
    if missing:
        raise FileRejected("INVALID_TEMPLATE", f"File không đúng mẫu. Thiếu cột: {', '.join(missing)}")
    idx = {c: header.index(c) for c in COLUMNS if c in header}
    body = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
    if len(body) > MAX_ROWS:
        raise FileRejected("TOO_MANY_ROWS", f"File có {len(body)} dòng, tối đa {MAX_ROWS}")
    res = ParseResult(total=len(body))
    seen: dict[tuple[str, date], int] = {}
    known = {s.lower() for s in known_skus} if known_skus is not None else None
    for i, r in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in r):
            continue

        def cell(col: str, r: list[Any] = r) -> Any:
            j = idx.get(col)
            return r[j] if j is not None and j < len(r) else None

        row_errors: list[RowError] = []
        sku_raw = cell("ma_sku")
        sku = str(sku_raw).strip() if sku_raw not in (None, "") else ""
        if not sku:
            row_errors.append(_err(i, "ma_sku", "REQUIRED"))
        cost_raw = cell("gia_von")
        cost: Decimal | None = None
        if cost_raw in (None, ""):
            row_errors.append(_err(i, "gia_von", "REQUIRED"))
        else:
            cost = _parse_cost(cost_raw)
            if cost is None:
                row_errors.append(_err(i, "gia_von", "INVALID_NUMBER"))
            elif cost < 0:
                row_errors.append(_err(i, "gia_von", "NEGATIVE"))
        date_raw = cell("ap_dung_tu")
        vf: date | None = None
        if date_raw in (None, ""):
            row_errors.append(_err(i, "ap_dung_tu", "REQUIRED"))
        else:
            vf = _parse_date(date_raw)
            if vf is None:
                row_errors.append(_err(i, "ap_dung_tu", "INVALID_DATE"))
        if not row_errors and sku and vf is not None:
            key = (sku.lower(), vf)
            if key in seen:
                row_errors.append(_err(i, "ma_sku", "DUPLICATE_IN_FILE"))
            else:
                seen[key] = i
        res.errors.extend(row_errors)
        if row_errors or cost is None or vf is None:
            continue
        if known is not None and sku.lower() not in known:
            res.errors.append(_err(i, "ma_sku", "UNKNOWN_SKU"))
        note_raw = cell("ghi_chu")
        note = str(note_raw).strip()[:500] if note_raw not in (None, "") else None
        res.rows.append(ParsedRow(i, sku, cost.quantize(Decimal("0.01")), vf, note))
    return res


def build_template(skus: list[tuple[str, str]]) -> bytes:
    """Mẫu Excel có sẵn danh sách SKU đang bán của khách để chỉ cần điền giá (spec 05 §6.4)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "gia_von"
    ws.append(list(COLUMNS))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for sku, name in skus:
        ws.append([sku, None, None, name])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 40
    dv = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, error="Giá vốn phải là số ≥ 0"
    )
    ws.add_data_validation(dv)
    dv.add(f"B2:B{max(2, len(skus) + 500)}")
    guide = wb.create_sheet("huong_dan")
    lines = [
        "HƯỚNG DẪN ĐIỀN GIÁ VỐN",
        "",
        "ma_sku: mã SKU trùng với SKU người bán (seller SKU) trên sàn. Bắt buộc.",
        "gia_von: giá vốn của một đơn vị, tính bằng VND, không âm. Chấp nhận 120000, 120.000 hoặc 120,000. Bắt buộc.",
        "ap_dung_tu: ngày bắt đầu áp dụng giá vốn, dạng dd/mm/yyyy. Bắt buộc. Đơn đặt từ ngày này dùng giá vốn mới.",
        "ghi_chu: ghi chú tùy ý. Không bắt buộc.",
        "",
        "Danh sách SKU đang bán đã được điền sẵn ở sheet gia_von (cột ghi_chu là tên sản phẩm, có thể xóa).",
        "Dòng nào lỗi sẽ được báo rõ dòng và lý do; các dòng hợp lệ vẫn được nhận.",
        "Mỗi file tối đa 5.000 dòng và 2 MB.",
        "Giá vốn mới sẽ được áp dụng sau lần cập nhật dữ liệu kế tiếp (khoảng 10–20 phút).",
    ]
    for line in lines:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
