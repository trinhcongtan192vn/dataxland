"""FR-INP-01, FR-INP-02, FR-INP-04 (spec 05 §6.3–6.5)."""

from __future__ import annotations

import io
from typing import Any

import openpyxl
import pytest
from sqlalchemy import func, select

from tests.conftest import auth


@pytest.fixture(autouse=True)
def _no_engine(monkeypatch: pytest.MonkeyPatch) -> None:
    # SKU đang bán lấy qua bộ máy chỉ số; test này không cần Cube.
    monkeypatch.setattr("app.routers.inputs._selling_skus", lambda ctx: [("DS-SRM-C30", "Serum"), ("DS-KEM-50", "Kem")])


def _xlsx(rows: list[list[Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "gia_von"
    ws.append(["ma_sku", "gia_von", "ap_dung_tu", "ghi_chu"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_fr_inp_02_excel_partial_accept(client: Any, session: Any) -> None:
    from app.db.models import ProductCost, UploadJob

    rows = [
        ["DS-SRM-C30", 110000, "01/10/2026", "ok"],
        ["DS-KEM-50", "95.000", "01/10/2026", None],
        ["DS-SUA-150", "52,000", "01/10/2026", None],  # SKU lạ: chỉ cảnh báo, vẫn nhận
        ["DS-TONER-200", 70000, "02/10/2026", None],
        ["DS-MASK-10", 45000, "03/10/2026", None],
        ["DS-SON-01", 60000, "04/10/2026", None],
        ["DS-CHONG-50", 98000, "05/10/2026", None],
        ["", 1000, "01/10/2026", None],  # dòng 9: thiếu SKU
        ["DS-X", "abc", "01/10/2026", None],  # dòng 10: số sai
        ["DS-Y", -5, "31/02/2026", None],  # dòng 11: âm và ngày sai
    ]
    r = client.post(
        "/v1/inputs/product-costs/upload",
        headers=auth("owner_seller"),
        files={"file": ("gia_von.xlsx", _xlsx(rows), "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows_total"] == 10
    assert body["rows_ok"] == 7
    assert body["rows_error"] == 3
    hard = [(e["row"], e["column"], e["code"]) for e in body["errors"] if not e["warning"]]
    assert (9, "ma_sku", "REQUIRED") in hard
    assert (10, "gia_von", "INVALID_NUMBER") in hard
    assert (11, "gia_von", "NEGATIVE") in hard
    assert (11, "ap_dung_tu", "INVALID_DATE") in hard
    assert any(e["code"] == "UNKNOWN_SKU" and e["row"] == 4 for e in body["errors"])
    assert all(e["message_vi"] for e in body["errors"])
    n = session.scalar(select(func.count()).select_from(ProductCost).where(ProductCost.tenant_id == "demo_seller"))
    assert n == 7
    job = session.scalar(select(UploadJob))
    assert job.rows_ok == 7 and job.rows_error == 3
    # Tải lại cùng file: upsert theo (seller_sku, ap_dung_tu), không tạo trùng.
    client.post(
        "/v1/inputs/product-costs/upload",
        headers=auth("owner_seller"),
        files={"file": ("gia_von.xlsx", _xlsx(rows), "application/octet-stream")},
    )
    session.expire_all()
    assert (
        session.scalar(select(func.count()).select_from(ProductCost).where(ProductCost.tenant_id == "demo_seller")) == 7
    )


def test_excel_duplicate_and_template(client: Any) -> None:
    rows = [["A1", 1000, "01/10/2026", None], ["a1", 2000, "01/10/2026", None]]
    body = client.post(
        "/v1/inputs/product-costs/upload",
        headers=auth("owner_seller"),
        files={"file": ("x.xlsx", _xlsx(rows), "application/octet-stream")},
    ).json()
    assert body["rows_ok"] == 1
    assert [(e["row"], e["code"]) for e in body["errors"] if not e["warning"]] == [(3, "DUPLICATE_IN_FILE")]
    r = client.post(
        "/v1/inputs/product-costs/upload",
        headers=auth("owner_seller"),
        files={"file": ("x.pdf", b"%PDF", "application/pdf")},
    )
    assert r.status_code == 400 and r.json()["code"] == "INVALID_FILE_TYPE"
    r = client.get("/v1/inputs/product-costs/template", headers=auth("owner_seller"))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["gia_von", "huong_dan"]
    assert [c.value for c in wb["gia_von"][1]] == ["ma_sku", "gia_von", "ap_dung_tu", "ghi_chu"]
    assert wb["gia_von"]["A2"].value == "DS-SRM-C30"


def test_fr_inp_01_cost_effective_dates_api(client: Any) -> None:
    h = auth("owner_seller")
    a = client.post(
        "/v1/inputs/product-costs",
        json={"seller_sku": "SKU1", "unit_cost_vnd": 100000, "valid_from": "2026-09-01"},
        headers=h,
    )
    assert a.status_code == 201
    b = client.post(
        "/v1/inputs/product-costs",
        json={"seller_sku": "SKU1", "unit_cost_vnd": 120000, "valid_from": "2026-10-01"},
        headers=h,
    )
    assert b.status_code == 201
    rows = client.get("/v1/inputs/product-costs?sku=SKU1", headers=h).json()
    assert [(r["valid_from"], r["valid_to"], r["unit_cost_vnd"]) for r in rows] == [
        ("2026-09-01", "2026-09-30", 100000),
        ("2026-10-01", None, 120000),
    ]
    on = client.get("/v1/inputs/product-costs?active_on=2026-09-15", headers=h).json()
    assert [r["unit_cost_vnd"] for r in on] == [100000]
    # Đổi valid_from thì tạo bản mới.
    r = client.put(f"/v1/inputs/product-costs/{b.json()['id']}", json={"valid_from": "2026-10-10"}, headers=h)
    assert r.status_code == 200 and r.json()["valid_from"] == "2026-10-10"
    assert len(client.get("/v1/inputs/product-costs?sku=SKU1", headers=h).json()) == 3
    assert client.delete(f"/v1/inputs/product-costs/{a.json()['id']}", headers=h).status_code == 204
    assert len(client.get("/v1/inputs/product-costs?sku=SKU1", headers=h).json()) == 2
    # Member chỉ đọc.
    assert client.get("/v1/inputs/product-costs", headers=auth("member_seller")).status_code == 200
    r = client.post(
        "/v1/inputs/product-costs",
        json={"seller_sku": "X", "unit_cost_vnd": 1, "valid_from": "2026-10-01"},
        headers=auth("member_seller"),
    )
    assert r.status_code == 403
    r = client.post(
        "/v1/inputs/product-costs", json={"seller_sku": "X", "unit_cost_vnd": -1, "valid_from": "2026-10-01"}, headers=h
    )
    assert r.status_code == 422


def test_fr_inp_04_budget_platform_validation(client: Any, session: Any) -> None:
    from dx_common.events.tables import events

    h = auth("owner_agency")
    r = client.post(
        "/v1/inputs/ad-budgets", json={"platform": "lazada_ads", "month": "2026-10", "budget_vnd": 1000}, headers=h
    )
    assert r.status_code == 400 and r.json()["code"] == "PLATFORM_NOT_IN_TENANT"
    r = client.post(
        "/v1/inputs/ad-budgets", json={"platform": "meta_ads", "month": "2026-10", "budget_vnd": 100000000}, headers=h
    )
    assert r.status_code == 201
    r2 = client.post(
        "/v1/inputs/ad-budgets", json={"platform": "meta_ads", "month": "2026-10", "budget_vnd": 120000000}, headers=h
    )
    assert r2.status_code == 201 and r2.json()["id"] == r.json()["id"] and r2.json()["budget_vnd"] == 120000000
    # Seller không có tiktok_ads.
    r = client.post(
        "/v1/inputs/ad-budgets",
        json={"platform": "tiktok_ads", "month": "2026-10", "budget_vnd": 1},
        headers=auth("owner_seller"),
    )
    assert r.status_code == 400
    # Partner được ghi ngân sách nhưng không được ghi giá vốn.
    assert client.put(
        f"/v1/inputs/ad-budgets/{r.json().get('id', 'x')}", json={"budget_vnd": 5}, headers=auth("partner")
    ).status_code in (403, 404)
    r = client.post(
        "/v1/inputs/ad-budgets",
        json={"platform": "google_ads", "month": "2026-10", "budget_vnd": 60000000},
        headers=auth("partner"),
    )
    assert r.status_code == 201
    assert client.get("/v1/inputs/product-costs", headers=auth("partner")).status_code == 403
    rows = client.get("/v1/inputs/ad-budgets?from_month=2026-10&to_month=2026-10", headers=h).json()
    assert {x["platform"] for x in rows} == {"meta_ads", "google_ads"}
    kinds = [r.payload["kind"] for r in session.execute(events.select().where(events.c.type == "inputs.changed"))]
    assert kinds.count("ad_budgets") == 3


def test_fr_ops_01_sync_tenants(tmp_path: Any, session: Any, client: Any) -> None:
    import shutil

    from app.registry_sync import sync_tenants
    from tests.conftest import REPO

    d = tmp_path / "tenants"
    shutil.copytree(REPO / "tenants", d)
    text = (d / "demo_seller.yml").read_text(encoding="utf-8").replace('sync: "0 */6 * * *"', 'sync: "0 */2 * * *"')
    (d / "demo_seller.yml").write_text(text, encoding="utf-8")
    (d / "broken01.yml").write_text("tenant_id: broken01\nname: X\npack: nope\n", encoding="utf-8")
    res = sync_tenants(session, d)
    session.commit()
    assert "demo_seller" in res.upserted
    assert any("broken01" in p for p in res.errors)
    cfg = client.get("/v1/tenants/demo_seller", headers=auth("owner_seller")).json()["config"]
    assert cfg["schedule"]["sync"] == "0 */2 * * *"
    assert all("secret_ref" not in s for s in cfg["sources"])
    # Khách khác không bị ảnh hưởng bởi file lỗi.
    assert client.get("/v1/tenants/demo_agency", headers=auth("owner_agency")).status_code == 200
    (d / "demo_agency.yml").unlink()
    res = sync_tenants(session, d)
    session.commit()
    assert res.archived == ["demo_agency"]
    assert client.get("/v1/tenants/demo_agency", headers=auth("data_lead")).status_code == 404


def test_fr_inp_05_campaign_budgets(client: Any, session: Any) -> None:
    h = auth("owner_agency")
    post = lambda body: client.post("/v1/inputs/ad-budgets", json=body, headers=h)  # noqa: E731
    r = post({"platform": "tiktok_ads", "month": "2026-11", "campaign_id": "tic001", "budget_vnd": 20000000})
    assert r.status_code == 201 and r.json()["campaign_id"] == "tic001"
    # Cùng chiến dịch, cùng tháng: cập nhật dòng cũ thay vì thêm dòng.
    r2 = post({"platform": "tiktok_ads", "month": "2026-11", "campaign_id": "tic001", "budget_vnd": 25000000})
    assert r2.status_code == 201 and r2.json()["id"] == r.json()["id"] and r2.json()["budget_vnd"] == 25000000
    assert (
        post({"platform": "tiktok_ads", "month": "2026-11", "campaign_id": "tic002", "budget_vnd": 1}).status_code
        == 201
    )
    # Đã nhập theo chiến dịch thì không nhập cả nền tảng cùng tháng (tránh cộng trùng), và ngược lại.
    c = post({"platform": "tiktok_ads", "month": "2026-11", "budget_vnd": 40000000})
    assert c.status_code == 409 and c.json()["code"] == "BUDGET_LEVEL_CONFLICT"
    assert post({"platform": "meta_ads", "month": "2026-11", "budget_vnd": 5}).status_code == 201
    c2 = post({"platform": "meta_ads", "month": "2026-11", "campaign_id": "m1", "budget_vnd": 5})
    assert c2.status_code == 409
    rows = client.get("/v1/inputs/ad-budgets?from_month=2026-11&to_month=2026-11", headers=h).json()
    assert [(b["platform"], b["campaign_id"]) for b in rows] == [
        ("meta_ads", None),
        ("tiktok_ads", "tic001"),
        ("tiktok_ads", "tic002"),
    ]


def test_fr_inp_05_campaign_list_via_metrics(client: Any) -> None:
    from app.metrics.cube_client import set_cube
    from tests.test_metrics import FakeCube

    fake = FakeCube()
    set_cube(fake)
    try:
        r = client.get("/v1/inputs/ad-campaigns", headers=auth("owner_agency"))
    finally:
        set_cube(None)
    assert r.status_code == 200, r.text
    assert r.json() and r.json()[0]["campaign_id"]
    assert all(t == "demo_agency" for t, _, _ in fake.calls)
