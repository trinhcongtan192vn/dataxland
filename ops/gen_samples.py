"""Sinh dữ liệu mẫu raw cho hai khách giả lập (spec 00 §5, spec 02 §8, spec 07 §10).

- Seed cố định: cùng `--anchor` cho kết quả giống hệt từng byte.
- `--anchor` là ngày cuối cùng có dữ liệu (theo giờ Việt Nam). Bản trong repo dùng 2026-10-15.
- Không có dữ liệu khách thật: tên, số điện thoại, email người mua đều là giả.

Đầu ra: `transform/samples/<source>/<object>.jsonl`, mỗi dòng là một bản ghi raw đúng hợp đồng spec 01 §1
(`_dx_loaded_at`, `_dx_source_account`, `_dx_record_hash`), và `transform/samples/files/` cho kiểm tra sức khỏe.

Cách dùng:
    uv run python ops/gen_samples.py                      # ghi vào transform/samples, anchor 2026-10-15
    uv run python ops/gen_samples.py --anchor 2026-11-02 --out .local/samples --loaded-at now
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import random
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

VN = ZoneInfo("Asia/Ho_Chi_Minh")
SEED = 20261016
DAYS = 90
DEFAULT_ANCHOR = date(2026, 10, 15)
REPO = Path(__file__).resolve().parents[1]

# Mã tài khoản nguồn của từng khách giả lập, khớp tenants/*.yml.
ACCOUNTS = {
    "demo_seller": {
        "shopee": "100001",
        "tiktok_shop": "7000001",
        "lazada": "300001",
        "meta_ads": "act_500001",
    },
    "demo_agency": {
        "shopee": "100002",
        "meta_ads": "act_600001",
        "tiktok_ads": "6900000001",
        "google_ads": "123-456-7890",
        "amazon_ads": "ENTITY1AMZ",
        "amazon_seller": "A1SELLERDEMO",
    },
}


@dataclass(frozen=True)
class Product:
    sku: str
    name: str
    price: int  # VND, hoặc cent USD với Amazon
    cost: int | None
    barcode: str | None = None
    category: str = "Làm đẹp"


SELLER_PRODUCTS = [
    Product("DS-SRM-C30", "Serum Vitamin C 30ml", 289000, 110000, "8938500000011"),
    Product("DS-KEM-50", "Kem dưỡng ẩm 50g", 245000, 95000, "8938500000028"),
    Product("DS-SUA-150", "Sữa rửa mặt dịu nhẹ 150ml", 159000, 52000, "8938500000035"),
    Product("DS-TONER-200", "Toner hoa hồng 200ml", 199000, 70000, "8938500000042"),
    Product("DS-MASK-10", "Mặt nạ giấy hộp 10 miếng", 129000, 45000, None),
    Product("DS-SON-01", "Son kem lì màu đỏ gạch", 179000, 60000, "8938500000059"),
    Product("DS-CHONG-50", "Kem chống nắng SPF50 50ml", 269000, 98000, "8938500000066"),
    Product("DS-TAYTB-01", "Tẩy tế bào chết 100g", 189000, 66000, None),
    Product("DS-DAU-100", "Dầu tẩy trang 100ml", 219000, 80000, "8938500000073"),
    # SKU lỗ sau phí (spec 00 §5).
    Product("DS-LOSS-01", "Combo quà tặng mini", 99000, 84000, None, "Quà tặng"),
    # SKU chưa có giá vốn (spec 00 §5).
    Product("DS-NOCOST-01", "Bông tẩy trang 300 miếng", 69000, None, None, "Phụ kiện"),
]
# Cùng sản phẩm với DS-SRM-C30 nhưng tên và SKU khác, chỉ bán trên TikTok Shop (dữ liệu cho spec 10).
TIKTOK_ONLY = [Product("TT-SERUMC-30ML", "Tinh chất Vitamin C 30ml chính hãng", 279000, 110000, None)]
# SKU đổi giá vốn giữa kỳ: DS-SRM-C30 từ 110.000 lên 125.000 (spec 02 §8).
COST_CHANGE = ("DS-SRM-C30", 125000, 45)  # (sku, giá mới, số ngày trước anchor)

AGENCY_PRODUCTS = [
    Product("BA-BOT-500", "Bột protein whey 500g", 590000, 260000, "8938600000010", "Thực phẩm bổ sung"),
    Product("BA-VIT-60", "Vitamin tổng hợp 60 viên", 350000, 120000, "8938600000027", "Thực phẩm bổ sung"),
    Product("BA-OMEGA-90", "Omega 3 90 viên", 420000, 150000, "8938600000034", "Thực phẩm bổ sung"),
]
AMAZON_PRICES_CENTS = {"BA-BOT-500": 2499, "BA-VIT-60": 1599, "BA-OMEGA-90": 1999}

FIRST_NAMES = [
    "An",
    "Bình",
    "Chi",
    "Dũng",
    "Hà",
    "Hùng",
    "Lan",
    "Linh",
    "Minh",
    "Nam",
    "Ngọc",
    "Phương",
    "Quân",
    "Thảo",
    "Trang",
    "Tuấn",
    "Vy",
    "Yến",
]
LAST_NAMES = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Vũ", "Đặng", "Bùi", "Đỗ"]
CITIES = [
    "Quận 1, TP. Hồ Chí Minh",
    "Cầu Giấy, Hà Nội",
    "Hải Châu, Đà Nẵng",
    "Ninh Kiều, Cần Thơ",
    "Thủ Đức, TP. Hồ Chí Minh",
    "Đống Đa, Hà Nội",
]


def _hash_record(rec: dict[str, Any]) -> str:
    body = {k: v for k, v in rec.items() if not k.startswith("_dx_")}
    return hashlib.sha256(json.dumps(body, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _epoch(dt: datetime) -> int:
    return int(dt.timestamp())


def _iso(dt: datetime) -> str:
    return dt.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def _laz_ts(dt: datetime) -> str:
    return dt.astimezone(VN).strftime("%Y-%m-%d %H:%M:%S +0700")


class Gen:
    def __init__(self, anchor: date, loaded_at: datetime) -> None:
        self.anchor = anchor
        self.start = anchor - timedelta(days=DAYS - 1)
        self.loaded_at = loaded_at
        self.rng = random.Random(SEED)
        self.tables: dict[str, list[dict[str, Any]]] = {}
        # Bể người mua dùng chung giữa các sàn của demo_seller, để có người mua trùng số điện thoại (spec 10).
        self.buyers = [self._buyer(i) for i in range(160)]

    # ---------------------------------------------------------------- tiện ích
    def _buyer(self, i: int) -> dict[str, str]:
        r = self.rng
        name = f"{r.choice(LAST_NAMES)} {r.choice(FIRST_NAMES)}"
        phone = "09" + "".join(str(r.randint(0, 9)) for _ in range(8))
        email = f"buyer{i:03d}@mail.example"
        return {
            "name": name,
            "phone": phone,
            "email": email,
            "address": f"{r.randint(1, 300)} Đường số {r.randint(1, 30)}, {r.choice(CITIES)}",
        }

    def emit(self, source: str, obj: str, account: str, rec: dict[str, Any], loaded_at: datetime | None = None) -> None:
        when = min(loaded_at or self.loaded_at, self.loaded_at)
        rec = dict(rec)
        rec["_dx_loaded_at"] = _iso(when)
        rec["_dx_source_account"] = account
        rec["_dx_record_hash"] = _hash_record(rec)
        self.tables.setdefault(f"{source}/{obj}", []).append(rec)

    def days(self) -> list[date]:
        return [self.start + timedelta(days=i) for i in range(DAYS)]

    def age(self, d: date) -> int:
        return (self.anchor - d).days

    def order_time(self, d: date) -> datetime:
        r = self.rng
        local = datetime(d.year, d.month, d.day, r.randint(7, 23), r.randint(0, 59), r.randint(0, 59), tzinfo=VN)
        return local.astimezone(UTC)

    def lifecycle(self, d: date) -> str:
        """Trạng thái chung theo tuổi đơn: `cancelled`, `returned`, `completed`, `delivered`, `shipping`, `confirmed`, `pending`."""
        r = self.rng
        age = self.age(d)
        x = r.random()
        if x < 0.06:
            return "cancelled"
        if age >= 8 and x < 0.11:
            return "returned"
        if age >= 8:
            return "completed"
        if age >= 4:
            return "delivered" if x < 0.5 else "shipping"
        if age >= 1:
            return "shipping" if x < 0.6 else "confirmed"
        return "confirmed" if x < 0.7 else "pending"

    def lines(self, catalog: list[Product], weights: list[float] | None = None) -> list[tuple[Product, int]]:
        r = self.rng
        n = r.choices([1, 2, 3], weights=[0.65, 0.27, 0.08])[0]
        picked = r.choices(catalog, weights=weights, k=n)
        seen: dict[str, tuple[Product, int]] = {}
        for p in picked:
            q = r.choices([1, 2], weights=[0.85, 0.15])[0]
            prev = seen.get(p.sku)
            seen[p.sku] = (p, q + (prev[1] if prev else 0))
        return list(seen.values())

    def daily_volume(self, d: date, base: float, trend_cut: float = 1.0) -> int:
        r = self.rng
        weekday_boost = 1.25 if d.weekday() >= 5 else 1.0
        mid_month = 1.4 if d.day in (9, 10, 15, 25) else 1.0  # ngày sale
        lam = base * weekday_boost * mid_month * trend_cut
        return max(1, int(round(r.gauss(lam, lam * 0.25))))

    def settle_after(self, d: date, status: str) -> datetime | None:
        """Thời điểm quyết toán; None nếu chưa quyết toán."""
        if status in ("cancelled", "pending", "confirmed", "shipping"):
            return None
        settle_day = d + timedelta(days=10 + self.rng.randint(0, 3))
        if settle_day > self.anchor:
            return None
        return datetime(settle_day.year, settle_day.month, settle_day.day, 9, 0, tzinfo=VN).astimezone(UTC)

    # ---------------------------------------------------------------- Shopee
    def shopee(self, tenant: str, catalog: list[Product], base: float, unknown_status: bool) -> None:
        acct = ACCOUNTS[tenant]["shopee"]
        r = self.rng
        status_raw = {
            "cancelled": "CANCELLED",
            "returned": "TO_RETURN",
            "completed": "COMPLETED",
            "delivered": "TO_CONFIRM_RECEIVE",
            "shipping": "SHIPPED",
            "confirmed": "READY_TO_SHIP",
            "pending": "UNPAID",
        }
        weights = [1.0] * len(catalog)
        seq = 0
        for d in self.days():
            totals_count = 0
            totals_gross = 0
            for _ in range(self.daily_volume(d, base)):
                seq += 1
                order_sn = f"{d:%y%m%d}{acct[-2:]}{seq:05d}SP"
                created = self.order_time(d)
                state = self.lifecycle(d)
                st = status_raw[state]
                if state == "cancelled" and r.random() < 0.3:
                    st = "IN_CANCEL"
                if unknown_status and self.age(d) in (1, 2) and seq % 7 == 0:
                    st = "SELLER_HOLD"  # trạng thái lạ chỉ có trong file ghi đè của demo_seller (spec 02 §8)
                buyer = r.choice(self.buyers)
                lines = self.lines(catalog, weights)
                gross = 0
                discounted_total = 0
                for i, (p, q) in enumerate(lines):
                    disc_price = int(round(p.price * r.choice([1.0, 1.0, 0.95, 0.9]) / 1000) * 1000)
                    gross += p.price * q
                    discounted_total += disc_price * q
                    self.emit(
                        "shopee",
                        "order_items",
                        acct,
                        {
                            "order_sn": order_sn,
                            "item_id": 1000 + catalog.index(p),
                            "model_id": 50000 + catalog.index(p) * 10 + i,
                            "item_name": p.name,
                            "item_sku": p.sku,
                            "model_sku": p.sku,
                            "model_quantity_purchased": q,
                            "model_original_price": p.price,
                            "model_discounted_price": disc_price,
                        },
                        created + timedelta(minutes=5),
                    )
                seller_voucher = r.choice([0, 0, 0, 10000, 15000])
                shopee_voucher = r.choice([0, 0, 10000, 20000])
                buyer_ship = r.choice([0, 15000, 15000, 20000])
                total = max(0, discounted_total - seller_voucher - shopee_voucher + buyer_ship)
                upd = created + timedelta(days=min(self.age(d), 9), hours=r.randint(1, 20))
                upd = min(upd, self.loaded_at - timedelta(hours=1))
                rec = {
                    "order_sn": order_sn,
                    "order_status": st,
                    "create_time": _epoch(created),
                    "update_time": _epoch(upd),
                    "pay_time": None if st == "UNPAID" else _epoch(created + timedelta(minutes=3)),
                    "currency": "VND",
                    "total_amount": total,
                    "order_original_amount": gross,
                    "order_discounted_amount": discounted_total,
                    "seller_voucher": seller_voucher,
                    "shopee_voucher": shopee_voucher,
                    "estimated_shipping_fee": 32000,
                    "buyer_paid_shipping_fee": buyer_ship,
                    "buyer_user_id": 900000 + self.buyers.index(buyer),
                    "buyer_username": "user" + buyer["phone"][-4:],
                    "recipient_name": buyer["name"],
                    "recipient_phone": buyer["phone"],
                    "recipient_full_address": buyer["address"],
                }
                # Đơn đổi trạng thái sau khi giao: bản cũ COMPLETED, bản mới TO_RETURN (FR-ING-07).
                if state == "returned":
                    old = dict(rec, order_status="COMPLETED", update_time=_epoch(upd - timedelta(days=2)))
                    self.emit("shopee", "orders", acct, old, upd - timedelta(days=2) + timedelta(hours=1))
                self.emit("shopee", "orders", acct, rec, upd + timedelta(hours=1))
                totals_count += 1
                totals_gross += gross

                settled = self.settle_after(d, state)
                if settled and state != "cancelled":
                    commission = int(round(discounted_total * 0.09))
                    service = int(round(discounted_total * 0.05))
                    transaction = int(round(total * 0.04))
                    actual_ship = r.randint(28, 38) * 1000
                    rebate = 15000
                    reverse = 22000 if state == "returned" else 0
                    escrow = (
                        discounted_total
                        - seller_voucher
                        - commission
                        - service
                        - transaction
                        - max(0, actual_ship - rebate - buyer_ship)
                        - reverse
                    )
                    self.emit(
                        "shopee",
                        "escrow",
                        acct,
                        {
                            "order_sn": order_sn,
                            "escrow_release_time": _epoch(settled),
                            "buyer_total_amount": total,
                            "original_price": gross,
                            "seller_discount": gross - discounted_total,
                            "voucher_from_seller": seller_voucher,
                            "voucher_from_shopee": shopee_voucher,
                            "commission_fee": commission,
                            "service_fee": service,
                            "seller_transaction_fee": transaction,
                            "actual_shipping_fee": actual_ship,
                            "shopee_shipping_rebate": rebate,
                            "buyer_paid_shipping_fee": buyer_ship,
                            "reverse_shipping_fee": reverse,
                            "escrow_amount": escrow,
                        },
                        settled + timedelta(hours=2),
                    )
                if state == "returned":
                    p, q = lines[0]
                    reason = r.choice(
                        ["CHANGE_MIND", "DAMAGED_PRODUCT", "WRONG_ITEM", "NOT_RECEIPT", "DIFFERENT_DESCRIPTION"]
                    )
                    ret_time = upd - timedelta(hours=3)
                    self.emit(
                        "shopee",
                        "returns",
                        acct,
                        {
                            "return_sn": "R" + order_sn,
                            "order_sn": order_sn,
                            "item_id": 1000 + catalog.index(p),
                            "model_id": 50000 + catalog.index(p) * 10,
                            "item_sku": p.sku,
                            "reason": reason,
                            "status": "COMPLETED",
                            "create_time": _epoch(ret_time),
                            "update_time": _epoch(upd),
                            "refund_amount": int(round(p.price * 0.95 / 1000) * 1000) * q,
                            "return_ship_fee": 22000,
                        },
                        upd + timedelta(hours=1),
                    )
            noise = 1.0 + (0.003 if d.day % 5 == 0 else 0.0)
            self.emit(
                "shopee",
                "daily_totals",
                acct,
                {
                    "date": d.isoformat(),
                    "order_count": totals_count,
                    "gross_amount": int(round(totals_gross * noise)),
                    "currency": "VND",
                },
                datetime(d.year, d.month, d.day, 23, 30, tzinfo=VN).astimezone(UTC) + timedelta(hours=2),
            )
        for i, p in enumerate(catalog):
            self.emit(
                "shopee",
                "products",
                acct,
                {
                    "item_id": 1000 + i,
                    "model_id": 50000 + i * 10,
                    "item_sku": p.sku,
                    "model_sku": p.sku,
                    "item_name": p.name,
                    "model_name": "Mặc định",
                    "gtin_code": p.barcode,
                    "category_id": 100017,
                    "category_name": p.category,
                },
                self.loaded_at - timedelta(days=1),
            )

    # ---------------------------------------------------------------- TikTok Shop
    def tiktok_shop(self, tenant: str, catalog: list[Product], base: float) -> None:
        acct = ACCOUNTS[tenant]["tiktok_shop"]
        r = self.rng
        status_raw = {
            "cancelled": "CANCELLED",
            "returned": "COMPLETED",
            "completed": "COMPLETED",
            "delivered": "DELIVERED",
            "shipping": "IN_TRANSIT",
            "confirmed": "AWAITING_SHIPMENT",
            "pending": "UNPAID",
        }
        seq = 0
        for d in self.days():
            cnt = 0
            gross_sum = 0
            for _ in range(self.daily_volume(d, base)):
                seq += 1
                oid = f"57{d:%y%m%d}{seq:06d}"
                created = self.order_time(d)
                state = self.lifecycle(d)
                buyer = r.choice(self.buyers[:120])
                lines = self.lines(catalog)
                gross = 0
                sub_total = 0
                seller_disc_total = 0
                line_no = 0
                for p, q in lines:
                    for _ in range(q):  # TikTok: mỗi đơn vị một dòng
                        line_no += 1
                        sale = int(round(p.price * r.choice([1.0, 0.92, 0.9]) / 1000) * 1000)
                        seller_disc = p.price - sale
                        gross += p.price
                        sub_total += sale
                        seller_disc_total += seller_disc
                        self.emit(
                            "tiktok_shop",
                            "order_items",
                            acct,
                            {
                                "order_id": oid,
                                "id": f"{oid}{line_no:02d}",
                                "sku_id": f"17{catalog.index(p):04d}",
                                "product_id": f"1729{catalog.index(p):04d}",
                                "seller_sku": p.sku,
                                "product_name": p.name,
                                "sku_name": "Mặc định",
                                "original_price": p.price,
                                "sale_price": sale,
                                "seller_discount": seller_disc,
                                "platform_discount": 0,
                                "quantity": 1,
                            },
                            created + timedelta(minutes=5),
                        )
                platform_disc = r.choice([0, 0, 10000, 15000])
                shipping_fee = r.choice([0, 18000, 22000])
                total = max(0, sub_total - platform_disc + shipping_fee)
                upd = min(
                    created + timedelta(days=min(self.age(d), 9), hours=r.randint(1, 20)),
                    self.loaded_at - timedelta(hours=1),
                )
                st = status_raw[state]
                rec = {
                    "id": oid,
                    "status": st,
                    "create_time": _epoch(created),
                    "update_time": _epoch(upd),
                    "paid_time": None if st == "UNPAID" else _epoch(created + timedelta(minutes=2)),
                    "currency": "VND",
                    "original_total_product_price": gross,
                    "seller_discount": seller_disc_total,
                    "platform_discount": platform_disc,
                    "sub_total": sub_total,
                    "shipping_fee": shipping_fee,
                    "total_amount": total,
                    "user_id": f"u{self.buyers.index(buyer):05d}",
                    "buyer_email": buyer["email"],
                    "recipient_name": buyer["name"],
                    "recipient_phone_number": "(+84)" + buyer["phone"][1:],
                    "recipient_full_address": buyer["address"],
                }
                self.emit("tiktok_shop", "orders", acct, rec, upd + timedelta(hours=1))
                cnt += 1
                gross_sum += gross
                settled = self.settle_after(d, state)
                if settled and state != "cancelled":
                    comm = int(round(sub_total * 0.06))
                    txn = int(round(total * 0.04))
                    ship_cost = r.randint(20, 30) * 1000 - shipping_fee
                    aff = int(round(sub_total * 0.08)) if r.random() < 0.3 else 0
                    self.emit(
                        "tiktok_shop",
                        "settlements",
                        acct,
                        {
                            "order_id": oid,
                            "transaction_id": f"ST{oid}",
                            "statement_time": _epoch(settled),
                            "currency": "VND",
                            "revenue_amount": sub_total,
                            "platform_commission_amount": comm,
                            "transaction_fee_amount": txn,
                            "shipping_cost_amount": max(0, ship_cost),
                            "affiliate_commission_amount": aff,
                            "settlement_amount": sub_total - comm - txn - max(0, ship_cost) - aff,
                        },
                        settled + timedelta(hours=2),
                    )
                if state == "returned":
                    p, _ = lines[0]
                    ret = upd - timedelta(hours=4)
                    self.emit(
                        "tiktok_shop",
                        "returns",
                        acct,
                        {
                            "return_id": f"RT{oid}",
                            "order_id": oid,
                            "line_item_id": f"{oid}01",
                            "return_type": "RETURN_AND_REFUND",
                            "return_status": "RETURN_OR_REFUND_REQUEST_COMPLETE",
                            "return_reason": r.choice(["wrong_item", "damaged", "quality_issue", "no_longer_needed"]),
                            "create_time": _epoch(ret),
                            "update_time": _epoch(upd),
                            "refund_total": int(round(p.price * 0.92 / 1000) * 1000),
                            "return_shipping_fee": 20000,
                        },
                        upd + timedelta(hours=1),
                    )
            self.emit(
                "tiktok_shop",
                "daily_totals",
                acct,
                {
                    "date": d.isoformat(),
                    "order_count": cnt,
                    "gross_amount": gross_sum,
                    "currency": "VND",
                },
                datetime(d.year, d.month, d.day, 23, 45, tzinfo=VN).astimezone(UTC) + timedelta(hours=2),
            )
        for i, p in enumerate(catalog):
            self.emit(
                "tiktok_shop",
                "products",
                acct,
                {
                    "product_id": f"1729{i:04d}",
                    "sku_id": f"17{i:04d}",
                    "seller_sku": p.sku,
                    "title": p.name,
                    "sku_name": "Mặc định",
                    "barcode": p.barcode,
                    "category_name": p.category,
                },
                self.loaded_at - timedelta(days=1),
            )

    # ---------------------------------------------------------------- Lazada
    def lazada(self, tenant: str, catalog: list[Product], base: float) -> None:
        acct = ACCOUNTS[tenant]["lazada"]
        r = self.rng
        status_raw = {
            "cancelled": "canceled",
            "returned": "returned",
            "completed": "confirmed",
            "delivered": "delivered",
            "shipping": "shipped",
            "confirmed": "ready_to_ship",
            "pending": "pending",
        }
        seq = 0
        for d in self.days():
            cnt = 0
            gross_sum = 0
            for _ in range(self.daily_volume(d, base)):
                seq += 1
                oid = f"8{d:%y%m%d}{seq:05d}"
                created = self.order_time(d)
                state = self.lifecycle(d)
                if state == "shipping" and r.random() < 0.08:
                    state_raw = "failed_delivery"
                else:
                    state_raw = status_raw[state]
                buyer = r.choice(self.buyers[40:])
                lines = self.lines(catalog)
                gross = 0
                paid_total = 0
                upd = min(
                    created + timedelta(days=min(self.age(d), 9), hours=r.randint(1, 20)),
                    self.loaded_at - timedelta(hours=1),
                )
                item_no = 0
                for p, q in lines:
                    for _ in range(q):
                        item_no += 1
                        paid = int(round(p.price * r.choice([1.0, 0.95]) / 1000) * 1000)
                        gross += p.price
                        paid_total += paid
                        self.emit(
                            "lazada",
                            "order_items",
                            acct,
                            {
                                "order_item_id": f"{oid}{item_no:02d}",
                                "order_id": oid,
                                "sku": p.sku,
                                "shop_sku": f"{2000 + catalog.index(p)}_VNAMZ",
                                "name": p.name,
                                "item_price": p.price,
                                "paid_price": paid,
                                "voucher_seller": p.price - paid,
                                "voucher_platform": 0,
                                "status": state_raw,
                                "created_at": _laz_ts(created),
                                "updated_at": _laz_ts(upd),
                            },
                            upd + timedelta(hours=1),
                        )
                voucher_platform = r.choice([0, 0, 10000])
                rec = {
                    "order_id": oid,
                    "statuses": state_raw,
                    "created_at": _laz_ts(created),
                    "updated_at": _laz_ts(upd),
                    "currency": "VND",
                    "price": gross,
                    "voucher_seller": gross - paid_total,
                    "voucher_platform": voucher_platform,
                    "shipping_fee": r.choice([0, 16500]),
                    "customer_first_name": buyer["name"],
                    "address_phone": buyer["phone"],
                    "address_full": buyer["address"],
                }
                self.emit("lazada", "orders", acct, rec, upd + timedelta(hours=1))
                cnt += 1
                gross_sum += gross
                settled = self.settle_after(d, state)
                if settled and state != "cancelled" and state_raw != "failed_delivery":
                    fees = [
                        ("Commission", -int(round(paid_total * 0.07))),
                        ("Payment Fee", -int(round(paid_total * 0.02))),
                        ("Shipping Fee Paid by Seller", -r.randint(8, 15) * 1000),
                        ("Item Price Credit", paid_total),
                    ]
                    for j, (fname, amount) in enumerate(fees):
                        self.emit(
                            "lazada",
                            "transactions",
                            acct,
                            {
                                "transaction_id": f"T{oid}{j}",
                                "order_no": oid,
                                "order_item_no": f"{oid}01",
                                "transaction_type": "Orders-Marketing Fees"
                                if fname == "Commission"
                                else "Orders-Sales",
                                "fee_name": fname,
                                "amount": amount,
                                "transaction_date": settled.astimezone(VN).date().isoformat(),
                            },
                            settled + timedelta(hours=2),
                        )
                if state == "returned":
                    p, _ = lines[0]
                    self.emit(
                        "lazada",
                        "returns",
                        acct,
                        {
                            "reverse_order_id": f"RV{oid}",
                            "order_id": oid,
                            "order_item_id": f"{oid}01",
                            "reason_text": r.choice(
                                ["Change of mind", "Damaged item", "Wrong item", "Not as described"]
                            ),
                            "status": "REFUND_SUCCESS",
                            "refund_amount": int(round(p.price * 0.95 / 1000) * 1000),
                            "created_at": _laz_ts(upd - timedelta(hours=5)),
                        },
                        upd + timedelta(hours=1),
                    )
            self.emit(
                "lazada",
                "daily_totals",
                acct,
                {
                    "date": d.isoformat(),
                    "order_count": cnt,
                    "gross_amount": gross_sum,
                    "currency": "VND",
                },
                datetime(d.year, d.month, d.day, 23, 50, tzinfo=VN).astimezone(UTC) + timedelta(hours=2),
            )
        for i, p in enumerate(catalog):
            self.emit(
                "lazada",
                "products",
                acct,
                {
                    "item_id": 2000 + i,
                    "sku_id": 30000 + i,
                    "seller_sku": p.sku,
                    "name": p.name,
                    "ean_code": p.barcode,
                    "primary_category": p.category,
                },
                self.loaded_at - timedelta(days=1),
            )

    # ---------------------------------------------------------------- Amazon Seller (Windsor)
    def amazon_seller(self, tenant: str, base: float) -> None:
        acct = ACCOUNTS[tenant]["amazon_seller"]
        r = self.rng
        catalog = AGENCY_PRODUCTS
        status_raw = {
            "cancelled": "Canceled",
            "returned": "Shipped",
            "completed": "Shipped",
            "delivered": "Shipped",
            "shipping": "Shipped",
            "confirmed": "Unshipped",
            "pending": "Pending",
        }
        seq = 0
        for d in self.days():
            cut = 0.8 if self.age(d) < 7 else 1.0
            cnt = 0
            gross_c = 0
            for _ in range(self.daily_volume(d, base, cut)):
                seq += 1
                oid = f"111-{seq:07d}-{(seq * 7919) % 10000000:07d}"
                created = self.order_time(d)
                state = self.lifecycle(d)
                lines = self.lines(catalog)
                total_c = 0
                for j, (p, q) in enumerate(lines):
                    price = AMAZON_PRICES_CENTS[p.sku]
                    promo = 200 if r.random() < 0.2 else 0
                    total_c += price * q - promo
                    self.emit(
                        "amazon_seller",
                        "order_items",
                        acct,
                        {
                            "amazon_order_id": oid,
                            "order_item_id": f"{seq:08d}{j}",
                            "seller_sku": p.sku,
                            "asin": f"B0DEMO{catalog.index(p):04d}",
                            "title": p.name,
                            "quantity_ordered": q,
                            "item_price_amount": f"{price * q / 100:.2f}",
                            "promotion_discount_amount": f"{promo / 100:.2f}",
                            "currency": "USD",
                        },
                        created + timedelta(hours=1),
                    )
                gross_items = sum(AMAZON_PRICES_CENTS[p.sku] * q for p, q in lines)
                upd = min(created + timedelta(days=min(self.age(d), 5), hours=2), self.loaded_at - timedelta(hours=1))
                self.emit(
                    "amazon_seller",
                    "orders",
                    acct,
                    {
                        "amazon_order_id": oid,
                        "purchase_date": _iso(created),
                        "last_update_date": _iso(upd),
                        "order_status": status_raw[state],
                        "order_total_amount": f"{total_c / 100:.2f}",
                        "order_total_currency": "USD",
                        "marketplace_id": "ATVPDKIKX0DER",
                        "buyer_email": f"{oid[-6:]}@marketplace.amazon.example",
                    },
                    upd + timedelta(hours=1),
                )
                cnt += 1
                gross_c += gross_items
                settled = self.settle_after(d, state)
                if settled and state != "cancelled":
                    for k, (fee_type, rate) in enumerate([("Commission", 0.15), ("FBAPerUnitFulfillmentFee", 0.12)]):
                        self.emit(
                            "amazon_seller",
                            "financial_events",
                            acct,
                            {
                                "event_id": f"FE{seq:07d}{k}",
                                "amazon_order_id": oid,
                                "posted_date": _iso(settled),
                                "event_type": "ShipmentEvent",
                                "fee_type": fee_type,
                                "amount": f"{-total_c * rate / 100:.2f}",
                                "currency": "USD",
                            },
                            settled + timedelta(hours=3),
                        )
                    if state == "returned":
                        self.emit(
                            "amazon_seller",
                            "financial_events",
                            acct,
                            {
                                "event_id": f"FE{seq:07d}9",
                                "amazon_order_id": oid,
                                "posted_date": _iso(settled),
                                "event_type": "RefundEvent",
                                "fee_type": "RefundPrincipal",
                                "amount": f"{-total_c / 100:.2f}",
                                "currency": "USD",
                            },
                            settled + timedelta(hours=3),
                        )
            self.emit(
                "amazon_seller",
                "daily_totals",
                acct,
                {
                    "date": d.isoformat(),
                    "order_count": cnt,
                    "gross_amount": f"{gross_c / 100:.2f}",
                    "currency": "USD",
                },
                datetime(d.year, d.month, d.day, 23, 55, tzinfo=VN).astimezone(UTC) + timedelta(hours=2),
            )

    # ---------------------------------------------------------------- Quảng cáo (Windsor)
    def ads(
        self,
        source: str,
        tenant: str,
        daily_spend: float,
        platform_roas: float,
        currency: str = "VND",
        october_boost: float = 1.0,
    ) -> None:
        acct = ACCOUNTS[tenant][source]
        r = self.rng
        campaigns = [("Chuyển đổi – Bán chạy", 0.45), ("Remarketing 30 ngày", 0.30), ("Khách mới – Sở thích", 0.25)]
        for d in self.days():
            boost = october_boost if (d.month == self.anchor.month and d.year == self.anchor.year) else 1.0
            last7 = self.age(d) < 7
            for ci, (cname, share) in enumerate(campaigns):
                for ai in range(2):
                    spend = daily_spend * share * 0.5 * boost * r.uniform(0.85, 1.15)
                    if last7 and tenant == "demo_agency":
                        spend *= 1.15
                    roas = platform_roas * r.uniform(0.8, 1.2) * (1.15 if ci == 1 else 1.0)
                    if currency == "USD":
                        spend_v = round(spend, 2)
                        value = round(spend * roas, 2)
                        spend_out: Any = f"{spend_v:.2f}"
                        value_out: Any = f"{value:.2f}"
                        cpm_unit = 0.004
                    else:
                        spend_v = round(spend)
                        value = round(spend * roas)
                        spend_out = int(spend_v)
                        value_out = int(value)
                        cpm_unit = 90.0
                    impressions = int(spend_v / cpm_unit * r.uniform(0.9, 1.1)) if spend_v else 0
                    clicks = int(impressions * r.uniform(0.008, 0.02))
                    conversions = max(0, int(clicks * r.uniform(0.02, 0.06)))
                    self.emit(
                        source,
                        "ad_insights_daily",
                        acct,
                        {
                            "date": d.isoformat(),
                            "account_id": acct,
                            "campaign_id": f"{source[:2]}c{ci + 1:03d}",
                            "campaign_name": cname,
                            "adgroup_id": f"{source[:2]}a{ci + 1:03d}{ai + 1}",
                            "adgroup_name": f"Nhóm {ai + 1}",
                            "spend": spend_out,
                            "currency": currency,
                            "impressions": impressions,
                            "clicks": clicks,
                            "conversions": conversions,
                            "conversion_value": value_out,
                        },
                        datetime(d.year, d.month, d.day, 23, 59, tzinfo=VN).astimezone(UTC) + timedelta(hours=4),
                    )

    # ---------------------------------------------------------------- Portal và tỷ giá
    def portal(self) -> None:
        start = self.start - timedelta(days=30)
        created = datetime(start.year, start.month, start.day, 3, 0, tzinfo=UTC)
        n = 0
        for tenant, catalog in (("demo_seller", SELLER_PRODUCTS + TIKTOK_ONLY), ("demo_agency", AGENCY_PRODUCTS)):
            for p in catalog:
                if p.cost is None:
                    continue
                n += 1
                self.emit(
                    "portal",
                    "product_costs",
                    "portal",
                    {
                        "id": f"pc{n:04d}",
                        "tenant_id": tenant,
                        "seller_sku": p.sku,
                        "unit_cost_vnd": f"{p.cost:.2f}",
                        "valid_from": start.isoformat(),
                        "note": None,
                        "deleted_at": None,
                        "created_at": _iso(created),
                        "updated_at": _iso(created),
                    },
                )
            if tenant == "demo_seller":
                sku, new_cost, ago = COST_CHANGE
                n += 1
                vf = self.anchor - timedelta(days=ago)
                ts = datetime(vf.year, vf.month, vf.day, 2, 0, tzinfo=UTC)
                self.emit(
                    "portal",
                    "product_costs",
                    "portal",
                    {
                        "id": f"pc{n:04d}",
                        "tenant_id": tenant,
                        "seller_sku": sku,
                        "unit_cost_vnd": f"{new_cost:.2f}",
                        "valid_from": vf.isoformat(),
                        "note": "Nhà cung cấp tăng giá",
                        "deleted_at": None,
                        "created_at": _iso(ts),
                        "updated_at": _iso(ts),
                    },
                )
        budgets = {
            "demo_seller": {"meta_ads": 45_000_000},
            "demo_agency": {
                "meta_ads": 100_000_000,
                "tiktok_ads": 80_000_000,
                "google_ads": 60_000_000,
                "amazon_ads": 65_000_000,
            },
        }
        months = sorted({d.replace(day=1) for d in self.days()})
        m = 0
        for tenant, plats in budgets.items():
            for month in months:
                for plat, amount in plats.items():
                    m += 1
                    self.emit(
                        "portal",
                        "ad_budgets",
                        "portal",
                        {
                            "id": f"ab{m:04d}",
                            "tenant_id": tenant,
                            "platform": plat,
                            "month": month.isoformat(),
                            "budget_vnd": f"{amount:.2f}",
                            "deleted_at": None,
                            "created_at": _iso(created),
                            "updated_at": _iso(created),
                        },
                    )

    def fx(self) -> None:
        r = self.rng
        rate = 25950.0
        d = self.start - timedelta(days=14)
        while d <= self.anchor:
            rate += r.uniform(-25, 30)
            if d.weekday() < 5:  # Vietcombank không công bố cuối tuần trong dữ liệu mẫu
                self.emit(
                    "fx_vcb",
                    "fx_rates",
                    "vcb",
                    {
                        "date": d.isoformat(),
                        "currency": "USD",
                        "rate_vnd": f"{rate:.2f}",
                        "source": "vcb",
                    },
                    datetime(d.year, d.month, d.day, 11, 30, tzinfo=UTC),
                )
            d += timedelta(days=1)

    # ---------------------------------------------------------------- chạy
    def run(self) -> None:
        self.shopee("demo_seller", SELLER_PRODUCTS, 6.0, unknown_status=True)
        self.tiktok_shop("demo_seller", SELLER_PRODUCTS[:9] + TIKTOK_ONLY + SELLER_PRODUCTS[9:], 4.0)
        self.lazada("demo_seller", SELLER_PRODUCTS, 3.2)
        self.ads("meta_ads", "demo_seller", 1_500_000, 3.1)
        self.shopee("demo_agency", AGENCY_PRODUCTS, 3.0, unknown_status=False)
        self.amazon_seller("demo_agency", 3.0)
        self.ads("meta_ads", "demo_agency", 3_100_000, 4.6, october_boost=1.45)
        self.ads("tiktok_ads", "demo_agency", 2_500_000, 2.6)
        self.ads("google_ads", "demo_agency", 1_900_000, 2.3)
        self.ads("amazon_ads", "demo_agency", 80.0, 3.0, currency="USD")
        self.portal()
        self.fx()


# ---------------------------------------------------------------------- file cho kiểm tra sức khỏe
def build_hc_files(tables: dict[str, list[dict[str, Any]]], out: Path) -> None:
    """File xuất mẫu từ dữ liệu `demo_seller` (spec 07 §10), có cột thông tin người mua giả để kiểm tra loại bỏ."""
    import openpyxl

    files_dir = out / "files"
    files_dir.mkdir(parents=True, exist_ok=True)
    acct = ACCOUNTS["demo_seller"]

    def latest(rows: list[dict[str, Any]], key: str) -> list[dict[str, Any]]:
        best: dict[str, dict[str, Any]] = {}
        for row in rows:
            k = str(row[key])
            if k not in best or row["_dx_loaded_at"] >= best[k]["_dx_loaded_at"]:
                best[k] = row
        return [best[k] for k in sorted(best)]

    def of(table: str, account: str) -> list[dict[str, Any]]:
        return [x for x in tables.get(table, []) if x["_dx_source_account"] == account]

    def vn_dt(epoch: int) -> str:
        return datetime.fromtimestamp(epoch, tz=UTC).astimezone(VN).strftime("%Y-%m-%d %H:%M")

    def write_xlsx(name: str, sheet: str, header: list[str], rows: list[list[Any]]) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append(header)
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        # Bỏ dấu thời gian trong file zip để file sinh lại giống hệt.
        (files_dir / name).write_bytes(_normalize_xlsx(buf.getvalue()))

    def write_csv(name: str, header: list[str], rows: list[list[Any]]) -> None:
        sio = io.StringIO()
        w = csv.writer(sio, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
        (files_dir / name).write_bytes(sio.getvalue().encode("utf-8-sig"))

    sp_status = {
        "COMPLETED": "Hoàn thành",
        "TO_RETURN": "Trả hàng/Hoàn tiền",
        "CANCELLED": "Đã hủy",
        "IN_CANCEL": "Đang hủy",
        "SHIPPED": "Đang giao",
        "TO_CONFIRM_RECEIVE": "Đã giao",
        "READY_TO_SHIP": "Chờ lấy hàng",
        "UNPAID": "Chờ xác nhận",
        "SELLER_HOLD": "Chờ xác nhận",
    }
    orders = {o["order_sn"]: o for o in latest(of("shopee/orders", acct["shopee"]), "order_sn")}
    rows = []
    for it in of("shopee/order_items", acct["shopee"]):
        o = orders[it["order_sn"]]
        rows.append(
            [
                o["order_sn"],
                vn_dt(o["create_time"]),
                sp_status[o["order_status"]],
                it["model_sku"],
                it["item_name"],
                it["model_original_price"],
                it["model_discounted_price"],
                it["model_quantity_purchased"],
                o["order_original_amount"],
                o["seller_voucher"],
                o["shopee_voucher"],
                o["buyer_paid_shipping_fee"],
                o["recipient_name"],
                o["recipient_phone"],
                o["recipient_full_address"],
                o["buyer_username"],
            ]
        )
    write_xlsx(
        "shopee_orders.xlsx",
        "orders",
        [
            "Mã đơn hàng",
            "Ngày đặt hàng",
            "Trạng Thái Đơn Hàng",
            "SKU phân loại hàng",
            "Tên sản phẩm",
            "Giá gốc",
            "Giá ưu đãi",
            "Số lượng",
            "Tổng giá trị đơn hàng (VND)",
            "Mã giảm giá của Shop",
            "Mã giảm giá của Shopee",
            "Phí vận chuyển (dự kiến)",
            "Tên Người nhận",
            "Số điện thoại",
            "Địa chỉ nhận hàng",
            "Người Mua",
        ],
        rows,
    )

    esc = of("shopee/escrow", acct["shopee"])
    write_xlsx(
        "shopee_income.xlsx",
        "Income",
        [
            "Mã đơn hàng",
            "Ngày hoàn thành thanh toán",
            "Giá sản phẩm",
            "Phí cố định",
            "Phí Dịch Vụ",
            "Phí thanh toán",
            "Phí vận chuyển thực tế",
            "Phí vận chuyển được trợ giá từ Shopee",
            "Phí vận chuyển Người mua trả",
            "Phí trả hàng",
            "Tổng tiền đã thanh toán",
        ],
        [
            [
                e["order_sn"],
                datetime.fromtimestamp(e["escrow_release_time"], tz=UTC).astimezone(VN).strftime("%Y-%m-%d"),
                e["original_price"] - e["seller_discount"],
                e["commission_fee"],
                e["service_fee"],
                e["seller_transaction_fee"],
                e["actual_shipping_fee"],
                e["shopee_shipping_rebate"],
                e["buyer_paid_shipping_fee"],
                e["reverse_shipping_fee"],
                e["escrow_amount"],
            ]
            for e in esc
        ],
    )

    tk_orders = {o["id"]: o for o in of("tiktok_shop/orders", acct["tiktok_shop"])}
    tk_rows = []
    for it in of("tiktok_shop/order_items", acct["tiktok_shop"]):
        o = tk_orders[it["order_id"]]
        tk_rows.append(
            [
                o["id"],
                o["status"],
                it["seller_sku"],
                it["product_name"],
                1,
                it["original_price"],
                it["seller_discount"],
                o["platform_discount"],
                o["shipping_fee"],
                vn_dt(o["create_time"]),
                o["recipient_name"],
                o["recipient_phone_number"],
                o["recipient_full_address"],
                o["buyer_email"],
            ]
        )
    write_csv(
        "tiktok_shop_orders.csv",
        [
            "Order ID",
            "Order Status",
            "Seller SKU",
            "Product Name",
            "Quantity",
            "SKU Unit Original Price",
            "SKU Seller Discount",
            "SKU Platform Discount",
            "Shipping Fee After Discount",
            "Created Time",
            "Recipient",
            "Phone #",
            "Detail Address",
            "Buyer Email",
        ],
        tk_rows,
    )
    st = of("tiktok_shop/settlements", acct["tiktok_shop"])
    write_csv(
        "tiktok_shop_settlement.csv",
        [
            "Order/adjustment ID",
            "Order settled time",
            "Total Revenue",
            "Platform commission fee",
            "Transaction fee",
            "Shipping cost",
            "Affiliate Commission",
            "Total settlement amount",
        ],
        [
            [
                s["order_id"],
                datetime.fromtimestamp(s["statement_time"], tz=UTC).astimezone(VN).strftime("%Y/%m/%d"),
                s["revenue_amount"],
                -s["platform_commission_amount"],
                -s["transaction_fee_amount"],
                -s["shipping_cost_amount"],
                -s["affiliate_commission_amount"],
                s["settlement_amount"],
            ]
            for s in st
        ],
    )

    lz_orders = {o["order_id"]: o for o in of("lazada/orders", acct["lazada"])}
    lz_rows = []
    for it in of("lazada/order_items", acct["lazada"]):
        o = lz_orders[it["order_id"]]
        lz_rows.append(
            [
                it["order_item_id"],
                o["order_id"],
                it["sku"],
                it["name"],
                it["item_price"],
                it["paid_price"],
                it["voucher_seller"],
                it["status"],
                o["created_at"][:16],
                o["customer_first_name"],
                o["address_phone"],
                o["address_full"],
            ]
        )
    write_xlsx(
        "lazada_orders.xlsx",
        "sheet1",
        [
            "Order Item Id",
            "Order Number",
            "Seller SKU",
            "Item Name",
            "Unit Price",
            "Paid Price",
            "Seller Discount Total",
            "Status",
            "Created at",
            "Customer Name",
            "Shipping Phone Number",
            "Shipping Address",
        ],
        lz_rows,
    )
    lt = of("lazada/transactions", acct["lazada"])
    write_csv(
        "lazada_transactions.csv",
        ["Transaction Date", "Transaction Type", "Fee Name", "Amount", "Order No.", "Order Item No."],
        [
            [
                t["transaction_date"],
                t["transaction_type"],
                t["fee_name"],
                t["amount"],
                t["order_no"],
                t["order_item_no"],
            ]
            for t in lt
        ],
    )

    meta = of("meta_ads/ad_insights_daily", acct["meta_ads"])
    write_csv(
        "meta_ads.csv",
        [
            "Day",
            "Campaign name",
            "Campaign ID",
            "Ad set name",
            "Ad set ID",
            "Amount spent (VND)",
            "Impressions",
            "Link clicks",
            "Results",
            "Purchases conversion value",
        ],
        [
            [
                a["date"],
                a["campaign_name"],
                a["campaign_id"],
                a["adgroup_name"],
                a["adgroup_id"],
                a["spend"],
                a["impressions"],
                a["clicks"],
                a["conversions"],
                a["conversion_value"],
            ]
            for a in meta
        ],
    )

    # Quảng cáo TikTok và Google cho kiểm tra sức khỏe lấy từ demo_agency (demo_seller không có hai nguồn này).
    tka = of("tiktok_ads/ad_insights_daily", ACCOUNTS["demo_agency"]["tiktok_ads"])
    write_csv(
        "tiktok_ads.csv",
        [
            "Date",
            "Campaign name",
            "Campaign ID",
            "Ad group name",
            "Ad group ID",
            "Cost",
            "Impressions",
            "Clicks",
            "Conversions",
            "Total purchase value",
        ],
        [
            [
                a["date"],
                a["campaign_name"],
                a["campaign_id"],
                a["adgroup_name"],
                a["adgroup_id"],
                a["spend"],
                a["impressions"],
                a["clicks"],
                a["conversions"],
                a["conversion_value"],
            ]
            for a in tka
        ],
    )
    gga = of("google_ads/ad_insights_daily", ACCOUNTS["demo_agency"]["google_ads"])
    write_csv(
        "google_ads.csv",
        [
            "Day",
            "Campaign",
            "Campaign ID",
            "Ad group",
            "Ad group ID",
            "Cost",
            "Impr.",
            "Clicks",
            "Conversions",
            "Conv. value",
        ],
        [
            [
                a["date"],
                a["campaign_name"],
                a["campaign_id"],
                a["adgroup_name"],
                a["adgroup_id"],
                a["spend"],
                a["impressions"],
                a["clicks"],
                a["conversions"],
                a["conversion_value"],
            ]
            for a in gga
        ],
    )

    costs = [c for c in tables.get("portal/product_costs", []) if c["tenant_id"] == "demo_seller"]
    write_xlsx(
        "product_costs.xlsx",
        "gia_von",
        ["ma_sku", "gia_von", "ap_dung_tu", "ghi_chu"],
        [
            [
                c["seller_sku"],
                int(float(c["unit_cost_vnd"])),
                date.fromisoformat(c["valid_from"]).strftime("%d/%m/%Y"),
                c["note"] or "",
            ]
            for c in costs
        ],
    )


def _normalize_xlsx(data: bytes) -> bytes:
    """Ghi lại file zip với dấu thời gian cố định để kết quả sinh lại giống hệt."""
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re

                content = re.sub(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>", b"", content)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, content)
    return buf.getvalue()


def write_tables(tables: dict[str, list[dict[str, Any]]], out: Path) -> None:
    for key in sorted(tables):
        path = out / f"{key}.jsonl"
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8", newline="\n") as fh:
            for rec in tables[key]:
                fh.write(json.dumps(rec, ensure_ascii=False, sort_keys=True) + "\n")


def generate(anchor: date = DEFAULT_ANCHOR, loaded_at: datetime | None = None) -> dict[str, list[dict[str, Any]]]:
    la = loaded_at or datetime(anchor.year, anchor.month, anchor.day, 23, 0, tzinfo=VN).astimezone(UTC) + timedelta(
        hours=3
    )
    g = Gen(anchor, la)
    g.run()
    return g.tables


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument(
        "--anchor", default=DEFAULT_ANCHOR.isoformat(), help="ngày cuối có dữ liệu, YYYY-MM-DD hoặc 'yesterday'"
    )
    ap.add_argument("--out", default=str(REPO / "transform" / "samples"))
    ap.add_argument("--loaded-at", default=None, help="thời điểm nạp (ISO UTC) hoặc 'now'")
    ap.add_argument("--no-files", action="store_true", help="không sinh file kiểm tra sức khỏe")
    args = ap.parse_args()
    if args.anchor == "yesterday":
        anchor = datetime.now(VN).date() - timedelta(days=1)
    else:
        anchor = date.fromisoformat(args.anchor)
    loaded_at = None
    if args.loaded_at == "now":
        loaded_at = datetime.now(UTC).replace(microsecond=0)
    elif args.loaded_at:
        loaded_at = datetime.fromisoformat(args.loaded_at.replace("Z", "+00:00"))
    tables = generate(anchor, loaded_at)
    out = Path(args.out)
    write_tables(tables, out)
    if not args.no_files:
        build_hc_files(tables, out)
    total = sum(len(v) for v in tables.values())
    print(f"Đã sinh {total} bản ghi trong {len(tables)} bảng vào {out} (anchor {anchor})")


if __name__ == "__main__":
    main()
